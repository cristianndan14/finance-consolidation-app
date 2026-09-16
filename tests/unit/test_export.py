import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.repositories.export import ExportRow
from app.web.exports import _rows_to_csv_bytes, _rows_to_xlsx_bytes

_ROW = ExportRow(
    posted_date=date(2026, 3, 5),
    description="SUPERMERCADO XYZ",
    merchant_name="Supermercado XYZ",
    category_name="Supermercado",
    amount=Decimal("-12345.67"),
    currency="ARS",
    kind="purchase",
    issuer_name="Visa",
    last4="1234",
)


def test_csv_header_when_no_rows() -> None:
    body = _rows_to_csv_bytes([])
    text = body.decode("utf-8-sig")
    lines = text.strip().splitlines()

    assert len(lines) == 1
    assert lines[0].startswith("Fecha;Descripcion")


def test_csv_includes_row_data() -> None:
    body = _rows_to_csv_bytes([_ROW])
    text = body.decode("utf-8-sig")
    lines = text.strip().splitlines()

    assert len(lines) == 2
    assert "2026-03-05" in lines[1]
    assert "Supermercado XYZ" in lines[1]
    assert "-12345,67" in lines[1]
    assert "Visa ...1234" in lines[1]


def test_xlsx_header_when_no_rows() -> None:
    body = _rows_to_xlsx_bytes([])
    workbook = load_workbook(filename=io.BytesIO(body))
    sheet = workbook.active

    assert sheet.max_row == 1
    assert sheet.cell(row=1, column=1).value == "Fecha"


def test_xlsx_amount_is_numeric() -> None:
    body = _rows_to_xlsx_bytes([_ROW])
    workbook = load_workbook(filename=io.BytesIO(body))
    sheet = workbook.active

    assert sheet.max_row == 2
    amount_cell = sheet.cell(row=2, column=5)
    assert isinstance(amount_cell.value, float)
    assert amount_cell.value == -12345.67
