"""Export de transacciones confirmadas a CSV/XLSX.

Las rutas solo arman el archivo a partir de `ExportRow`; el filtrado por
usuario lo hace RLS via `db.user_tx`, no hay logica de seguridad extra aca.
El armado del archivo esta separado de la consulta (`_rows_to_csv_bytes` /
`_rows_to_xlsx_bytes` reciben una lista de `ExportRow` ya en memoria) para
poder testearlo sin Postgres.
"""

from __future__ import annotations

import csv
import io

from fastapi import APIRouter
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from starlette.responses import Response

from app.deps import CurrentUserDep
from app.infra import db
from app.repositories import analytics, export

router = APIRouter(tags=["export"])

_HEADERS = (
    "Fecha",
    "Descripcion",
    "Comercio",
    "Categoria",
    "Monto",
    "Moneda",
    "Tipo",
    "Tarjeta",
)


@router.get("/export/transactions.csv", include_in_schema=False)
async def export_csv(
    user: CurrentUserDep, year: int | None = None, month: int | None = None
) -> Response:
    rows, period = await _load(user, year, month)
    body = _rows_to_csv_bytes(rows)
    return Response(
        body,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{_filename(period, "csv")}"'},
    )


@router.get("/export/transactions.xlsx", include_in_schema=False)
async def export_xlsx(
    user: CurrentUserDep, year: int | None = None, month: int | None = None
) -> Response:
    rows, period = await _load(user, year, month)
    body = _rows_to_xlsx_bytes(rows)
    return Response(
        body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename(period, "xlsx")}"'},
    )


# ─────────────────────────────────────────────────────────────────────────────
# Interno
# ─────────────────────────────────────────────────────────────────────────────
async def _load(
    user: CurrentUserDep, year: int | None, month: int | None
) -> tuple[list[export.ExportRow], analytics.Period | None]:
    async with db.user_tx(user.db_claims) as conn:
        period = await _resolve_period(conn, year, month)
        if period is None:
            return [], None
        rows = await export.transactions_for_period(conn, year=period.year, month=period.month)
    return rows, period


async def _resolve_period(
    conn: db.AsyncConnection, year: int | None, month: int | None
) -> analytics.Period | None:
    if year and month and 1 <= month <= 12:
        return analytics.Period(year=year, month=month)
    latest = await analytics.available_periods(conn, limit=1)
    return latest[0] if latest else None


def _filename(period: analytics.Period | None, extension: str) -> str:
    if period is None:
        return f"transacciones.{extension}"
    return f"transacciones-{period.year:04d}-{period.month:02d}.{extension}"


# Caracteres que Excel/LibreOffice interpretan como el inicio de una formula.
# `description`/`merchant_name` vienen del texto del PDF (via el LLM), no del
# codigo: una linea que empiece asi ejecutaria como formula al abrir el export.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(value: str) -> str:
    if value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _row_values(row: export.ExportRow) -> tuple[str, str, str, str, str, str, str, str]:
    card = f"{row.issuer_name} ...{row.last4}" if row.issuer_name and row.last4 else ""
    return (
        row.posted_date.isoformat(),
        _csv_safe(row.description),
        _csv_safe(row.merchant_name),
        _csv_safe(row.category_name),
        str(row.amount),
        row.currency,
        row.kind,
        _csv_safe(card),
    )


def _rows_to_csv_bytes(rows: list[export.ExportRow]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_HEADERS)
    for row in rows:
        writer.writerow(_row_values(row))
    return buffer.getvalue().encode("utf-8-sig")


def _rows_to_xlsx_bytes(rows: list[export.ExportRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook() no genero una hoja activa")
    sheet.title = "Transacciones"
    sheet.append(_HEADERS)

    for row in rows:
        card = f"{row.issuer_name} ...{row.last4}" if row.issuer_name and row.last4 else ""
        sheet.append(
            (
                row.posted_date,
                _csv_safe(row.description),
                _csv_safe(row.merchant_name),
                _csv_safe(row.category_name),
                row.amount,
                row.currency,
                row.kind,
                _csv_safe(card),
            )
        )

    for column_index in range(1, len(_HEADERS) + 1):
        letter = get_column_letter(column_index)
        sheet.column_dimensions[letter].width = 18

    amount_column = get_column_letter(5)
    for cell in sheet[amount_column][1:]:
        cell.number_format = "#,##0.00"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
